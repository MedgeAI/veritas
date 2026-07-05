"""Window builder: reconstructs bounded evidence windows from source files.

Given an EvidenceLocator and the original source file, rebuilds the exact
bounded window (rows/cols range + highlights). The rebuild is deterministic:
same locator + same source file -> same window.

Fail-loud: if the source file is missing or hash mismatch, raises immediately.
Does not silently fall back to cached data.
"""

from __future__ import annotations

import csv
import hashlib
from pathlib import Path
from typing import Any

from engine.static_audit.evidence_windows.locator import EvidenceLocator
from engine.static_audit.evidence_windows.manifest import EvidenceWindowManifest


class EvidenceWindowError(Exception):
    """Raised when evidence window cannot be rebuilt."""


def build_window(
    source_root: Path,
    locator: EvidenceLocator,
) -> EvidenceWindowManifest:
    """Rebuild a bounded evidence window from the original source file.

    Args:
        source_root: Root directory containing source data files.
        locator: The evidence locator specifying the bounded region.

    Returns:
        EvidenceWindowManifest with the rebuilt window data.

    Raises:
        EvidenceWindowError: If the source file is missing or hash mismatch.
    """
    source_path = source_root / locator.source_path
    if not source_path.exists():
        raise EvidenceWindowError(
            f"Source file not found: {locator.source_path!r} "
            f"(resolved: {source_path}). Cannot rebuild evidence window."
        )

    # Verify file hash — fail loud on mismatch
    actual_sha256 = _compute_sha256(source_path)
    if actual_sha256 != locator.source_sha256:
        raise EvidenceWindowError(
            f"Source file hash mismatch for {locator.source_path!r}: "
            f"expected {locator.source_sha256}, got {actual_sha256}. "
            f"File may have been modified since the signal was produced."
        )

    suffix = source_path.suffix.lower()
    if suffix == ".xlsx":
        headers, data = _read_xlsx_window(source_path, locator)
    elif suffix in (".csv", ".tsv"):
        headers, data = _read_csv_window(source_path, locator, suffix)
    else:
        raise EvidenceWindowError(
            f"Unsupported source file type: {suffix!r}. "
            f"Supported: .xlsx, .csv, .tsv"
        )

    return EvidenceWindowManifest(
        schema_version="1.0",
        signal_id=locator.signal_id,
        detector_id=locator.detector_id,
        source_path=locator.source_path,
        source_sha256=locator.source_sha256,
        sheet=locator.sheet,
        rows=locator.rows,
        cols=locator.cols,
        highlight_rows=list(locator.highlight_rows),
        highlight_cols=list(locator.highlight_cols),
        headers=headers,
        data=data,
        rebuild_status="ok",
        extraction_method=locator.extraction_method or "veritas.window_builder",
    )


def build_window_from_locator(
    source_root: Path,
    locator_dict: dict[str, Any],
) -> EvidenceWindowManifest:
    """Convenience: build window from a locator dict (e.g., from JSON).

    Raises:
        EvidenceWindowError: If the locator dict is invalid or rebuild fails.
    """
    locator = EvidenceLocator.from_dict(locator_dict)
    return build_window(source_root, locator)


def _compute_sha256(path: Path) -> str:
    """Compute SHA-256 hex digest of a file."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def _read_xlsx_window(
    path: Path,
    locator: EvidenceLocator,
) -> tuple[list[str], list[list[Any]]]:
    """Read bounded window from an XLSX file.

    Returns (headers, data) where data is list of rows.
    """
    try:
        import openpyxl
    except ImportError as e:
        raise EvidenceWindowError(
            "openpyxl is required to read .xlsx files"
        ) from e

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    # Use locator.sheet; if empty, use the first sheet
    sheet_name = locator.sheet or wb.sheetnames[0]
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise EvidenceWindowError(
            f"Sheet {sheet_name!r} not found in {path.name}. "
            f"Available sheets: {wb.sheetnames}"
        )

    ws = wb[sheet_name]

    row_start, row_end = locator.parse_row_range()
    col_range = locator.parse_col_range()

    # Convert letter cols to numeric if needed
    if isinstance(col_range[0], str):
        col_start = _letter_to_index(col_range[0])
        col_end = _letter_to_index(col_range[1])
    else:
        col_start, col_end = col_range

    # Read all rows into a list for indexed access
    all_rows: list[list[Any]] = []
    for row in ws.iter_rows(min_row=1, values_only=True):
        all_rows.append(list(row))

    wb.close()

    if not all_rows:
        return ([], [])

    # First row is headers
    full_headers = [str(c) if c is not None else "" for c in all_rows[0]]
    # Slice headers to col range (col_start/col_end are 1-based)
    headers = full_headers[col_start - 1 : col_end]

    # Slice data rows (row_start/row_end are 1-based, row 1 is header)
    data: list[list[Any]] = []
    for row_idx in range(row_start - 1, min(row_end, len(all_rows))):
        row = all_rows[row_idx]
        sliced = row[col_start - 1 : col_end]
        data.append(sliced)

    return (headers, data)


def _read_csv_window(
    path: Path,
    locator: EvidenceLocator,
    suffix: str,
) -> tuple[list[str], list[list[Any]]]:
    """Read bounded window from a CSV/TSV file.

    Returns (headers, data) where data is list of rows.
    """
    delimiter = "\t" if suffix == ".tsv" else ","
    with open(path, "r", newline="", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f, delimiter=delimiter)
        all_rows = list(reader)

    if not all_rows:
        return ([], [])

    row_start, row_end = locator.parse_row_range()
    col_range = locator.parse_col_range()

    if isinstance(col_range[0], str):
        col_start = _letter_to_index(col_range[0])
        col_end = _letter_to_index(col_range[1])
    else:
        col_start, col_end = col_range

    headers = all_rows[0][col_start - 1 : col_end]

    data: list[list[Any]] = []
    for row_idx in range(row_start - 1, min(row_end, len(all_rows))):
        row = all_rows[row_idx]
        # Pad row if shorter than expected
        padded = row + [""] * max(0, col_end - len(row))
        sliced = padded[col_start - 1 : col_end]
        data.append(sliced)

    return (headers, data)


def _letter_to_index(letter: str) -> int:
    """Convert Excel column letter(s) to 1-based index. A=1, B=2, ..., Z=26, AA=27."""
    letter = letter.upper()
    result = 0
    for ch in letter:
        result = result * 26 + (ord(ch) - ord("A") + 1)
    return result
