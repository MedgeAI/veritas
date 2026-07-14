"""Real TypedVerifier adapters — wrap the deterministic source-data detectors as B3-B5 verifiers.

The ablation harness (agent_harness) consumes `TypedVerifier = (claim, case) -> VerifierSignal|None`.
This module turns the existing numeric forensics (duplicate columns / fixed relationships /
row-offset reuse / paired-ratio / duplicate-row-vector / narrow paired-difference) into a real,
non-mock verifier that runs on the case's actual xlsx and grounds its signal to the sheet locus.

Scope: a claim is IN SCOPE when its locus resolves to an `.xlsx` workbook (from
`claim.evidence.workbook/sheet` or the `evidence_span` target). Out of scope -> returns None so
the harness falls back to the agent. This is the source-data (L3) verifier; method/visual/code
verifiers (L2/L4/L1) are separate adapters to add later.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from engine.benchmark.agent_harness import ArtifactRenderer, TypedVerifier, VerifierSignal
from engine.benchmark.schema import BenchmarkCase, ClaimInstance, parse_evidence_span
from engine.static_audit.tools.source_data_findings import (
    duplicate_column_findings,
    fixed_relationship_findings,
    parse_workbook_vectors,
)
from engine.static_audit.tools.source_data_pair_forensics import (
    PairForensicsParams,
    duplicate_row_vector_findings,
    paired_difference_spread_findings,
    paired_ratio_reuse_findings,
    row_offset_scalar_findings,
)

_PARAMS = PairForensicsParams(min_pairs=8, min_support=0.95, ratio_places=4, max_offset=100,
                              max_findings_per_category=200, min_duplicate_row_width=2)
_RISK_ORDER = {"high": 3, "medium": 2, "low": 1}
_RISK_CONF = {"high": 0.9, "medium": 0.65, "low": 0.45}


def _resolve_locus(claim: ClaimInstance) -> tuple[str | None, str | None]:
    """(workbook, sheet) for a claim, from its Evidence or the evidence_span target."""
    ev = claim.evidence
    if ev.workbook:
        return ev.workbook, ev.sheet
    if claim.evidence_span:
        _relation, _source, target = parse_evidence_span(claim.evidence_span)
        parts = target.split(":")
        if parts and parts[0].endswith(".xlsx"):
            return parts[0], (parts[1] if len(parts) > 1 else None)
    return None, None


def _sheet_findings(path: Path, sheet: str | None) -> list[dict[str, Any]]:
    """Run every source-data detector on `sheet` (or all sheets if None)."""
    findings: list[dict[str, Any]] = []
    for sv in parse_workbook_vectors(path):
        if sheet is not None and sv.sheet != sheet:
            continue
        findings += duplicate_column_findings(sv, 8, 0.95, 200)
        findings += fixed_relationship_findings(sv, 8, 0.95, 200)
        findings += row_offset_scalar_findings(sv, _PARAMS)
        findings += paired_ratio_reuse_findings(sv, _PARAMS)
        findings += duplicate_row_vector_findings(sv, _PARAMS)
        findings += paired_difference_spread_findings(sv, _PARAMS)
    return findings


def source_data_verifier(data_dir: str | Path) -> TypedVerifier:
    """A real typed verifier over the source-data numeric detectors, reading xlsx from `data_dir`."""
    data_dir = Path(data_dir)

    def verify(claim: ClaimInstance, _case: BenchmarkCase) -> VerifierSignal | None:
        workbook, sheet = _resolve_locus(claim)
        if not workbook or not str(workbook).endswith(".xlsx"):
            return None  # not a source-data locus -> out of scope
        path = data_dir / workbook
        if not path.exists():
            return None
        target = f"{workbook}:{sheet}" if sheet else workbook
        span = f"{claim.level}:{workbook}->{target}"
        findings = _sheet_findings(path, sheet)
        if not findings:
            return VerifierSignal(claim.relation, fired=False, evidence_span=span, confidence=0.85)
        best = max(findings, key=lambda f: _RISK_ORDER.get(str(f.get("risk_level")), 0))
        conf = _RISK_CONF.get(str(best.get("risk_level")), 0.6)
        return VerifierSignal(claim.relation, fired=True, evidence_span=span, confidence=conf)

    return verify


def _sheet_to_text(path: Path, sheet: str | None, *, max_rows: int, max_cols: int) -> str:
    """Render a worksheet as a compact, row-numbered text table so an agent can inspect the data."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        titles = [sheet] if sheet and sheet in wb.sheetnames else wb.sheetnames[:1]
        blocks = []
        for title in titles:
            ws = wb[title]
            lines = [f"# {path.name} / {title}  ({ws.max_row} rows x {ws.max_column} cols)"]
            for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if r > max_rows:
                    lines.append(f"... ({ws.max_row - max_rows} more rows omitted)")
                    break
                cells = ["" if v is None else str(v) for v in row[:max_cols]]
                lines.append(f"r{r}: " + ", ".join(cells))
            blocks.append("\n".join(lines))
        return "\n\n".join(blocks)
    finally:
        wb.close()


def source_data_renderer(data_dir: str | Path, *, max_rows: int = 90, max_cols: int = 12) -> ArtifactRenderer:
    """An ArtifactRenderer that loads the claim's source-data sheet content into the prompt.

    Lets a bare agent actually READ the numbers (spot duplicated rows / fixed ratios itself) rather
    than only see file names. Falls back to the artifact inventory when the locus is not a resolvable
    xlsx. Rows/cols are capped so wide sheets stay within a token budget.
    """
    data_dir = Path(data_dir)

    def render(claim: ClaimInstance, case: BenchmarkCase) -> str:
        workbook, sheet = _resolve_locus(claim)
        if workbook and str(workbook).endswith(".xlsx"):
            path = data_dir / workbook
            if path.exists():
                return _sheet_to_text(path, sheet, max_rows=max_rows, max_cols=max_cols)
        from engine.benchmark.agent_harness import render_artifacts
        return render_artifacts(case)

    return render
