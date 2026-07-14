"""Single-factor XLSX injection primitives for synthetic twins."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from .contracts import CLAIM_TYPES, CellChange, InjectionOperation


def _dump_log(path: Path | None, operation: InjectionOperation, seed: int) -> None:
    if path is None:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(operation.log_dict(seed=seed), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def inject_duplicate_columns(
    *,
    source_path: Path,
    output_path: Path,
    sheet_name: str,
    source_column: str,
    target_column: str,
    start_row: int,
    end_row: int,
    seed: int,
    log_path: Path | None = None,
) -> InjectionOperation:
    """Copy one numeric column into another over a row span."""
    wb = load_workbook(source_path)
    ws = wb[sheet_name]
    changes: list[CellChange] = []
    for row in range(start_row, end_row + 1):
        source_ref = f"{source_column}{row}"
        target_ref = f"{target_column}{row}"
        source_cell = ws[source_ref]
        target_cell = ws[target_ref]
        original = target_cell.value
        target_cell.value = source_cell.value
        target_cell.number_format = source_cell.number_format
        changes.append(CellChange(ref=target_ref, orig=original, new=target_cell.value))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()

    operation = InjectionOperation(
        injection_class="duplicate_columns",
        claim_type=CLAIM_TYPES["duplicate_columns"],
        workbook=source_path.name,
        sheet=sheet_name,
        target=f"{source_path.name} / {sheet_name}",
        injection={
            "seed": seed,
            "op": "duplicate_columns",
            "src_col": source_column,
            "dst_col": target_column,
            "rows": f"{start_row}-{end_row}",
        },
        cells=changes,
    )
    _dump_log(log_path, operation, seed)
    return operation


def annotations_yaml(operation: InjectionOperation, *, base_paper_id: str) -> str:
    description = (
        f"{operation.injection['op']} in {operation.target}: "
        f"{operation.injection}"
    )
    return "\n".join(
        [
            "paper:",
            f'  base_paper_id: "{base_paper_id}"',
            '  source: "injected"',
            "claims:",
            f'  - claim_type: "{operation.claim_type}"',
            f'    target: "{operation.target}"',
            f'    description: "{description}"',
            '    evidence_type: "source_data"',
            "    confirmed_by_human: false",
            "    injection:",
            *[f"      {key}: {json.dumps(value)}" for key, value in operation.injection.items()],
            "    deterministically_verifiable: true",
            "",
        ]
    )


def inject_duplicate_row_vector(
    *,
    source_path: Path,
    output_path: Path,
    sheet_name: str,
    source_row: int,
    target_row: int,
    columns: list[str],
    seed: int,
    log_path: Path | None = None,
) -> InjectionOperation:
    wb = load_workbook(source_path)
    ws = wb[sheet_name]
    changes: list[CellChange] = []
    for col in columns:
        src = ws[f"{col}{source_row}"]
        dst = ws[f"{col}{target_row}"]
        original = dst.value
        dst.value = src.value
        dst.number_format = src.number_format
        changes.append(CellChange(ref=f"{col}{target_row}", orig=original, new=dst.value))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    operation = InjectionOperation(
        injection_class="duplicate_row_vector",
        claim_type=CLAIM_TYPES["duplicate_row_vector"],
        workbook=source_path.name,
        sheet=sheet_name,
        target=f"{source_path.name} / {sheet_name}",
        injection={
            "seed": seed,
            "op": "duplicate_row_vector",
            "src_rows": str(source_row),
            "dst_rows": str(target_row),
            "columns": columns,
        },
        cells=changes,
    )
    _dump_log(log_path, operation, seed)
    return operation


def inject_fixed_ratio(
    *,
    source_path: Path,
    output_path: Path,
    sheet_name: str,
    source_column: str,
    target_column: str,
    start_row: int,
    end_row: int,
    factor: float,
    seed: int,
    log_path: Path | None = None,
) -> InjectionOperation:
    return _inject_column_transform(
        source_path=source_path,
        output_path=output_path,
        sheet_name=sheet_name,
        source_column=source_column,
        target_column=target_column,
        start_row=start_row,
        end_row=end_row,
        transform=lambda value: float(value) * factor,
        injection_class="fixed_ratio",
        injection_extra={"factor": factor},
        seed=seed,
        log_path=log_path,
    )


def inject_fixed_difference(
    *,
    source_path: Path,
    output_path: Path,
    sheet_name: str,
    source_column: str,
    target_column: str,
    start_row: int,
    end_row: int,
    delta: float,
    seed: int,
    log_path: Path | None = None,
) -> InjectionOperation:
    return _inject_column_transform(
        source_path=source_path,
        output_path=output_path,
        sheet_name=sheet_name,
        source_column=source_column,
        target_column=target_column,
        start_row=start_row,
        end_row=end_row,
        transform=lambda value: float(value) + delta,
        injection_class="fixed_difference",
        injection_extra={"delta": delta},
        seed=seed,
        log_path=log_path,
    )


def inject_row_offset_exact_reuse(
    *,
    source_path: Path,
    output_path: Path,
    sheet_name: str,
    column: str,
    start_row: int,
    end_row: int,
    offset: int,
    ratio: float,
    seed: int,
    log_path: Path | None = None,
) -> InjectionOperation:
    wb = load_workbook(source_path)
    ws = wb[sheet_name]
    changes: list[CellChange] = []
    for row in range(start_row, end_row + 1):
        src = ws[f"{column}{row}"]
        dst_ref = f"{column}{row + offset}"
        dst = ws[dst_ref]
        original = dst.value
        dst.value = float(src.value) * ratio
        dst.number_format = src.number_format
        changes.append(CellChange(ref=dst_ref, orig=original, new=dst.value))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    operation = InjectionOperation(
        injection_class="row_offset_exact_reuse",
        claim_type=CLAIM_TYPES["row_offset_exact_reuse"],
        workbook=source_path.name,
        sheet=sheet_name,
        target=f"{source_path.name} / {sheet_name}",
        injection={
            "seed": seed,
            "op": "row_offset_exact_reuse",
            "offset_k": offset,
            "ratio": ratio,
            "rows": f"{start_row}-{end_row}",
            "col": column,
        },
        cells=changes,
    )
    _dump_log(log_path, operation, seed)
    return operation


def inject_paired_difference_spread(
    *,
    source_path: Path,
    output_path: Path,
    sheet_name: str,
    left_column: str,
    right_column: str,
    start_row: int,
    end_row: int,
    band: float,
    seed: int,
    log_path: Path | None = None,
) -> InjectionOperation:
    wb = load_workbook(source_path)
    ws = wb[sheet_name]
    changes: list[CellChange] = []
    for row in range(start_row, end_row + 1):
        left = ws[f"{left_column}{row}"]
        right = ws[f"{right_column}{row}"]
        original = right.value
        sign = -1 if row % 2 else 1
        right.value = float(left.value) + sign * band
        right.number_format = left.number_format
        changes.append(CellChange(ref=f"{right_column}{row}", orig=original, new=right.value))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    operation = InjectionOperation(
        injection_class="paired_difference_spread",
        claim_type=CLAIM_TYPES["paired_difference_spread"],
        workbook=source_path.name,
        sheet=sheet_name,
        target=f"{source_path.name} / {sheet_name}",
        injection={
            "seed": seed,
            "op": "paired_difference_spread",
            "pair_cols": [left_column, right_column],
            "rows": f"{start_row}-{end_row}",
            "band": band,
        },
        cells=changes,
    )
    _dump_log(log_path, operation, seed)
    return operation


def inject_cross_sheet_duplication(
    *,
    source_path: Path,
    output_path: Path,
    source_sheet: str,
    target_sheet: str,
    source_column: str,
    target_column: str,
    start_row: int,
    end_row: int,
    seed: int,
    log_path: Path | None = None,
) -> InjectionOperation:
    wb = load_workbook(source_path)
    ws_src = wb[source_sheet]
    ws_dst = wb[target_sheet]
    changes: list[CellChange] = []
    for row in range(start_row, end_row + 1):
        src = ws_src[f"{source_column}{row}"]
        dst_ref = f"{target_sheet}!{target_column}{row}"
        dst = ws_dst[f"{target_column}{row}"]
        original = dst.value
        dst.value = src.value
        dst.number_format = src.number_format
        changes.append(CellChange(ref=dst_ref, orig=original, new=dst.value))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    operation = InjectionOperation(
        injection_class="cross_sheet_duplication",
        claim_type=CLAIM_TYPES["cross_sheet_duplication"],
        workbook=source_path.name,
        sheet=f"{source_sheet} -> {target_sheet}",
        target=f"{source_path.name} / {source_sheet} -> {target_sheet}",
        injection={
            "seed": seed,
            "op": "cross_sheet_duplication",
            "sheetA": source_sheet,
            "sheetB": target_sheet,
            "block": f"{source_column}{start_row}:{source_column}{end_row} -> {target_column}{start_row}:{target_column}{end_row}",
        },
        cells=changes,
    )
    _dump_log(log_path, operation, seed)
    return operation


def inject_p_value_inconsistency(
    *,
    source_path: Path,
    output_path: Path,
    sheet_name: str,
    p_column: str,
    row: int,
    multiplier: float,
    seed: int,
    log_path: Path | None = None,
) -> InjectionOperation:
    wb = load_workbook(source_path)
    ws = wb[sheet_name]
    cell = ws[f"{p_column}{row}"]
    original = cell.value
    cell.value = float(original) * multiplier
    change = CellChange(ref=f"{p_column}{row}", orig=original, new=cell.value)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    operation = InjectionOperation(
        injection_class="p_value_inconsistency",
        claim_type=CLAIM_TYPES["p_value_inconsistency"],
        workbook=source_path.name,
        sheet=sheet_name,
        target=f"{source_path.name} / {sheet_name}",
        injection={
            "seed": seed,
            "op": "p_value_inconsistency",
            "row": row,
            "col_P": p_column,
            "multiplier": multiplier,
        },
        cells=[change],
    )
    _dump_log(log_path, operation, seed)
    return operation


def inject_grim_violation(
    *,
    source_path: Path,
    output_path: Path,
    sheet_name: str,
    mean_column: str,
    row: int,
    new_value: float,
    n: int,
    seed: int,
    log_path: Path | None = None,
) -> InjectionOperation:
    wb = load_workbook(source_path)
    ws = wb[sheet_name]
    cell = ws[f"{mean_column}{row}"]
    original = cell.value
    cell.value = new_value
    change = CellChange(ref=f"{mean_column}{row}", orig=original, new=cell.value)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    operation = InjectionOperation(
        injection_class="grim_violation",
        claim_type=CLAIM_TYPES["grim_violation"],
        workbook=source_path.name,
        sheet=sheet_name,
        target=f"{source_path.name} / {sheet_name}",
        injection={
            "seed": seed,
            "op": "grim_violation",
            "row": row,
            "col_mean": mean_column,
            "n": n,
        },
        cells=[change],
    )
    _dump_log(log_path, operation, seed)
    return operation


def _inject_column_transform(
    *,
    source_path: Path,
    output_path: Path,
    sheet_name: str,
    source_column: str,
    target_column: str,
    start_row: int,
    end_row: int,
    transform,
    injection_class: str,
    injection_extra: dict,
    seed: int,
    log_path: Path | None,
) -> InjectionOperation:
    wb = load_workbook(source_path)
    ws = wb[sheet_name]
    changes: list[CellChange] = []
    for row in range(start_row, end_row + 1):
        src = ws[f"{source_column}{row}"]
        dst = ws[f"{target_column}{row}"]
        original = dst.value
        dst.value = transform(src.value)
        dst.number_format = src.number_format
        changes.append(CellChange(ref=f"{target_column}{row}", orig=original, new=dst.value))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    operation = InjectionOperation(
        injection_class=injection_class,
        claim_type=CLAIM_TYPES[injection_class],
        workbook=source_path.name,
        sheet=sheet_name,
        target=f"{source_path.name} / {sheet_name}",
        injection={
            "seed": seed,
            "op": injection_class,
            "src_col": source_column,
            "dst_col": target_column,
            "rows": f"{start_row}-{end_row}",
            **injection_extra,
        },
        cells=changes,
    )
    _dump_log(log_path, operation, seed)
    return operation
