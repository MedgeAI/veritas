"""Workbook table discovery for injection substrates."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from zipfile import BadZipFile

from openpyxl import load_workbook


NUMERIC_SUMMARY_TERMS = (
    "estimate",
    "se",
    "std",
    "t",
    "z",
    "p",
    "p-value",
    "p value",
    "mean",
    "sd",
    "hr",
    "ci",
    "loghr",
    "log hr",
)

GENE_LIST_TERMS = ("gene", "symbol", "ensembl", "transcript")


@dataclass(frozen=True)
class ParsedTable:
    workbook: str
    sheet: str
    header_row: int
    headers: list[str]
    numeric_columns: list[int]


def _stringify_row(values: Iterable[object]) -> list[str]:
    return ["" if value is None else str(value).strip() for value in values]


def find_header_row(path: Path, sheet_name: str, *, max_scan_rows: int = 8) -> int | None:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    best_row: int | None = None
    best_score = 0
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_scan_rows, values_only=True),
        start=1,
    ):
        labels = _stringify_row(row)
        text = " ".join(labels).lower()
        score = sum(1 for term in NUMERIC_SUMMARY_TERMS if term in text)
        nonempty = sum(1 for label in labels if label)
        if score > best_score and nonempty >= 2:
            best_row = row_idx
            best_score = score
    wb.close()
    return best_row


def discover_numeric_summary_tables(path: Path) -> list[ParsedTable]:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, OSError, ValueError):
        return []
    tables: list[ParsedTable] = []
    for ws in wb.worksheets:
        header_row = find_header_row(path, ws.title)
        if header_row is None:
            continue
        headers = _stringify_row(
            next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
        )
        header_text = " ".join(headers).lower()
        if any(term in header_text for term in GENE_LIST_TERMS) and not any(
            term in header_text for term in ("estimate", "se", "p", "mean", "sd", "hr")
        ):
            continue
        numeric_columns: list[int] = []
        for col_idx, _header in enumerate(headers, start=1):
            seen = 0
            numeric = 0
            for row in ws.iter_rows(
                min_row=header_row + 1,
                max_row=min(ws.max_row, header_row + 25),
                min_col=col_idx,
                max_col=col_idx,
                values_only=True,
            ):
                value = row[0]
                if value in (None, ""):
                    continue
                seen += 1
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    numeric += 1
            if seen and numeric / seen >= 0.5:
                numeric_columns.append(col_idx)
        if len(numeric_columns) >= 2:
            tables.append(
                ParsedTable(
                    workbook=path.name,
                    sheet=ws.title,
                    header_row=header_row,
                    headers=headers,
                    numeric_columns=numeric_columns,
                )
            )
    wb.close()
    return tables
