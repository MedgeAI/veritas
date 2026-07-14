"""Matched-clean controls for FAR (false-accusation-rate) measurement.

A dirty twin proves the auditor CAN catch an injected inconsistency. It says nothing about how
often the auditor cries wolf on genuine data. FAR needs the other half: benign loci that
structurally resemble a fabrication but are real. This module mints them from the PRISTINE
mother workbook (a genuine, non-retracted paper): every pair of real-numeric columns is a
"could these two columns be a duplicate / ratio / fixed-difference of each other?" locus, and in
authentic data the honest answer is no. A detector that flags one is making a false accusation.

Design choices that keep the ground truth honest:
  * Column universe = the columns the DETECTOR itself sees as numeric (`parse_workbook_vectors`
    -> numeric_columns with >= min_overlap depth), so the clean GT lives in the same space where
    FAR is defined — not an injection-side view that might disagree (this paper stores some
    numbers as text; those columns are invisible to the detector and correctly excluded).
  * Injected pairs are excluded (they are dirty in the corresponding twin).
  * Each control is materialised as a label=clean claim row (see `clean_control_rows`) so it
    round-trips through engine.benchmark.schema.load_annotations and engine.benchmark.metrics.

FAR is then simply: of these clean loci, the fraction a system flags. `flagged_pairs` reads a
system's findings; `far` divides. Run it before and after the whitelist to show triage earns the
low FAR by discrimination, not by abstaining.
"""

from __future__ import annotations

from dataclasses import dataclass
from itertools import combinations
from pathlib import Path

from openpyxl.utils import get_column_letter

from engine.static_audit.tools.source_data_findings import (
    duplicate_column_findings,
    fixed_relationship_findings,
    parse_workbook_vectors,
)


@dataclass(frozen=True)
class CleanControl:
    """One benign column-pair locus in a pristine mother table."""

    workbook: str
    sheet: str
    col_a: str
    col_b: str

    @property
    def pair_key(self) -> frozenset[str]:
        return frozenset((self.col_a, self.col_b))

    @property
    def target(self) -> str:
        return f"{self.workbook} / {self.sheet} cols {self.col_a},{self.col_b} (benign control)"


def _numeric_columns(sheet_vectors, min_overlap: int) -> list[int]:
    return sorted(
        col for col, cells in sheet_vectors.numeric_columns.items() if len(cells) >= min_overlap
    )


def enumerate_clean_controls(
    pristine_workbook: Path,
    sheet: str,
    *,
    min_overlap: int = 8,
    exclude_pairs: set[frozenset[str]] | None = None,
) -> list[CleanControl]:
    """All real-numeric column pairs in `sheet` of the pristine mother, minus injected pairs.

    Degenerate/invisible columns are handled by construction: `parse_workbook_vectors` only
    reports genuinely-numeric columns, and `min_overlap` drops columns too sparse to pair.
    """
    exclude_pairs = exclude_pairs or set()
    controls: list[CleanControl] = []
    for sv in parse_workbook_vectors(pristine_workbook):
        if sv.sheet != sheet:
            continue
        cols = _numeric_columns(sv, min_overlap)
        for a, b in combinations(cols, 2):
            la, lb = get_column_letter(a), get_column_letter(b)
            if frozenset((la, lb)) in exclude_pairs:
                continue
            controls.append(CleanControl(pristine_workbook.name, sheet, la, lb))
    return controls


def flagged_pairs(findings: list[dict]) -> set[frozenset[str]]:
    """Column pairs a system flagged (from finding `column_pair` fields)."""
    out: set[frozenset[str]] = set()
    for f in findings:
        pair = f.get("column_pair")
        if pair and len(pair) == 2:
            out.add(frozenset(str(p) for p in pair))
    return out


def detect_pairs(pristine_workbook: Path, sheet: str, *, min_overlap: int = 8, min_support: float = 0.95, limit: int = 500) -> list[dict]:
    """Run the numeric duplicate/relationship detectors on one sheet of the pristine mother."""
    findings: list[dict] = []
    for sv in parse_workbook_vectors(pristine_workbook):
        if sv.sheet != sheet:
            continue
        findings.extend(duplicate_column_findings(sv, min_overlap, min_support, limit))
        findings.extend(fixed_relationship_findings(sv, min_overlap, min_support, limit))
    return findings


def far(controls: list[CleanControl], flagged: set[frozenset[str]]) -> dict[str, float]:
    """False-accusation rate over the clean controls: flagged clean loci / total clean loci."""
    total = len(controls)
    if total == 0:
        return {"far": 0.0, "flagged": 0, "total": 0}
    hit = sum(1 for c in controls if c.pair_key in flagged)
    return {"far": hit / total, "flagged": hit, "total": total}


def clean_control_rows(controls: list[CleanControl], *, claim_type: str = "source_data.duplicate_columns") -> list[dict]:
    """Render clean controls as annotation rows (label=clean) for an annotations.yaml `claims` list.

    claim_type names the *question* the locus poses to the detector; the label makes it a control.
    """
    rows: list[dict] = []
    for i, c in enumerate(controls):
        rows.append(
            {
                "claim_type": claim_type,
                "label": "clean",
                "target": c.target,
                "description": f"benign real-data column pair {c.col_a}/{c.col_b}; genuine independent measurements",
                "evidence_type": "source_data",
                "confirmed_by_human": False,
                "deterministically_verifiable": True,
            }
        )
    return rows
