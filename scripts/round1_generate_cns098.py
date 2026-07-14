from __future__ import annotations

import json
import shutil
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from engine.static_audit.tools.source_data_findings import (
    duplicate_column_findings,
    fixed_relationship_findings,
    parse_workbook_vectors,
)
from engine.static_audit.tools.source_data_pair_forensics import (
    PairForensicsParams,
    analyze_xlsx_root,
)
from engine.static_audit.stats.statcheck import audit_t_test_rows
from engine.twins.contracts import CLAIM_TYPES, CellChange, InjectionOperation
from engine.twins.injector import (
    annotations_yaml,
    inject_duplicate_columns,
    inject_duplicate_row_vector,
    inject_fixed_difference,
    inject_fixed_ratio,
    inject_p_value_inconsistency,
    inject_paired_difference_spread,
    inject_row_offset_exact_reuse,
)
from engine.twins.table_parser import discover_numeric_summary_tables


PAPER_ID = "CNS-098_898e4634_epigenetic_clocks_174_diseases"


def _copy_supplementary(base_dir: Path, twin_dir: Path) -> Path:
    supplementary = twin_dir / "supplementary"
    if supplementary.exists():
        shutil.rmtree(supplementary)
    supplementary.mkdir(parents=True)
    for workbook in sorted(base_dir.glob("*.xlsx")):
        shutil.copy2(workbook, supplementary / workbook.name)
    return supplementary


def _write_annotation(twin_dir: Path, operation: InjectionOperation) -> None:
    (twin_dir / "annotations.yaml").write_text(
        annotations_yaml(operation, base_paper_id=PAPER_ID),
        encoding="utf-8",
    )


def _write_log(twin_dir: Path, operation: InjectionOperation, seed: int) -> None:
    (twin_dir / "injection_log.json").write_text(
        json.dumps(operation.log_dict(seed=seed), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _source_data_cleanliness(base_dir: Path) -> dict:
    errors = []
    findings = []
    for workbook_path in sorted(base_dir.glob("*.xlsx")):
        try:
            sheets = parse_workbook_vectors(workbook_path)
        except Exception as exc:
            errors.append({"workbook": workbook_path.name, "error": type(exc).__name__})
            continue
        for sheet in sheets:
            findings.extend(duplicate_column_findings(sheet, 12, 0.98, 200))
            findings.extend(fixed_relationship_findings(sheet, 12, 0.98, 200))
    pair = analyze_xlsx_root(
        base_dir,
        PairForensicsParams(
            min_pairs=8,
            min_support=0.95,
            ratio_places=4,
            max_offset=80,
            max_findings_per_category=50,
            min_duplicate_row_width=2,
        ),
    )
    categories = Counter(item["category"] for item in findings + pair["findings"])
    return {
        "source_data_findings": len(findings),
        "pair_forensics_findings": len(pair["findings"]),
        "priority_pair_findings": len(pair["priority_findings"]),
        "categories": dict(categories),
        "errors": errors + pair["errors"],
    }


def _statcheck_probe(base_dir: Path) -> list[dict]:
    path = base_dir / "41467_2025_66106_MOESM6_ESM.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["s5"]
    rows = []
    for row_idx in range(2, ws.max_row + 1):
        rows.append(
            {
                "row": row_idx,
                "t": ws[f"E{row_idx}"].value,
                "df": ws[f"D{row_idx}"].value,
                "P": ws[f"F{row_idx}"].value,
            }
        )
    wb.close()
    return audit_t_test_rows(rows, t_column="t", df_column="df", p_column="P")


def _inject_cross_workbook(
    *,
    supplementary: Path,
    source_name: str,
    target_name: str,
    source_sheet: str,
    target_sheet: str,
    source_column: str,
    target_column: str,
    start_row: int,
    end_row: int,
    seed: int,
) -> InjectionOperation:
    src_wb = load_workbook(supplementary / source_name, read_only=True, data_only=True)
    src_ws = src_wb[source_sheet]
    dst_wb = load_workbook(supplementary / target_name)
    dst_ws = dst_wb[target_sheet]
    changes = []
    for row in range(start_row, end_row + 1):
        new_value = src_ws[f"{source_column}{row}"].value
        dst = dst_ws[f"{target_column}{row}"]
        original = dst.value
        dst.value = new_value
        changes.append(
            CellChange(
                ref=f"{target_name}!{target_sheet}!{target_column}{row}",
                orig=original,
                new=new_value,
            )
        )
    src_wb.close()
    dst_wb.save(supplementary / target_name)
    dst_wb.close()
    return InjectionOperation(
        injection_class="cross_sheet_duplication",
        claim_type=CLAIM_TYPES["cross_sheet_duplication"],
        workbook=f"{source_name} -> {target_name}",
        sheet=f"{source_sheet} -> {target_sheet}",
        target=f"{source_name}/{source_sheet} -> {target_name}/{target_sheet}",
        injection={
            "seed": seed,
            "op": "cross_sheet_duplication",
            "sheetA": f"{source_name}/{source_sheet}",
            "sheetB": f"{target_name}/{target_sheet}",
            "block": f"{source_column}{start_row}:{source_column}{end_row} -> {target_column}{start_row}:{target_column}{end_row}",
        },
        cells=changes,
    )


def _generate_one(base_dir: Path, twins_root: Path, klass: str, seq: int) -> dict:
    seed = 98000 + seq
    twin_dir = twins_root / f"{PAPER_ID}__{klass}__{seq:02d}"
    twin_dir.mkdir(parents=True)
    supplementary = _copy_supplementary(base_dir, twin_dir)

    long_table = supplementary / "41467_2025_66106_MOESM2_ESM.xlsx"
    stat_table = supplementary / "41467_2025_66106_MOESM6_ESM.xlsx"
    start = 3 + (seq - 1) * 20
    all_start = 3
    all_end = 176
    paired_cols = [("E", "F"), ("G", "H"), ("I", "J")]
    left_col, right_col = paired_cols[seq - 1]

    if klass == "duplicate_columns":
        operation = inject_duplicate_columns(
            source_path=long_table,
            output_path=long_table,
            sheet_name="s1",
            source_column=left_col,
            target_column=right_col,
            start_row=all_start,
            end_row=all_end,
            seed=seed,
        )
    elif klass == "duplicate_row_vector":
        operation = inject_duplicate_row_vector(
            source_path=long_table,
            output_path=long_table,
            sheet_name="s1",
            source_row=start,
            target_row=start + 80,
            columns=["C", "D", "E", "F", "G", "H"],
            seed=seed,
        )
    elif klass == "fixed_ratio":
        operation = inject_fixed_ratio(
            source_path=long_table,
            output_path=long_table,
            sheet_name="s1",
            source_column=left_col,
            target_column=right_col,
            start_row=all_start,
            end_row=all_end,
            factor=2.0,
            seed=seed,
        )
    elif klass == "fixed_difference":
        operation = inject_fixed_difference(
            source_path=long_table,
            output_path=long_table,
            sheet_name="s1",
            source_column=left_col,
            target_column=right_col,
            start_row=all_start,
            end_row=all_end,
            delta=0.5,
            seed=seed,
        )
    elif klass == "row_offset_exact_reuse":
        operation = inject_row_offset_exact_reuse(
            source_path=long_table,
            output_path=long_table,
            sheet_name="s1",
            column="E",
            start_row=3,
            end_row=89,
            offset=87,
            ratio=1.0,
            seed=seed,
        )
    elif klass == "paired_difference_spread":
        operation = inject_paired_difference_spread(
            source_path=long_table,
            output_path=long_table,
            sheet_name="s1",
            left_column=left_col,
            right_column=right_col,
            start_row=all_start,
            end_row=all_end,
            band=0.01,
            seed=seed,
        )
    elif klass == "cross_sheet_duplication":
        operation = _inject_cross_workbook(
            supplementary=supplementary,
            source_name="41467_2025_66106_MOESM7_ESM.xlsx",
            target_name="41467_2025_66106_MOESM5_ESM.xlsx",
            source_sheet="s6",
            target_sheet="s4",
            source_column="D",
            target_column="B",
            start_row=2,
            end_row=14,
            seed=seed,
        )
    elif klass == "p_value_inconsistency":
        operation = inject_p_value_inconsistency(
            source_path=stat_table,
            output_path=stat_table,
            sheet_name="s5",
            p_column="F",
            row=2 + seq,
            multiplier=10.0,
            seed=seed,
        )
    else:
        raise ValueError(klass)

    _write_log(twin_dir, operation, seed)
    _write_annotation(twin_dir, operation)
    return {
        "twin": str(twin_dir),
        "class": klass,
        "seq": seq,
        "changed_cells": len(operation.cells),
    }


def main() -> int:
    root = Path("input/twins_work/CNS-098")
    twins_root = root / "twins"
    if twins_root.exists():
        shutil.rmtree(twins_root)
    twins_root.mkdir(parents=True)
    classes = [
        "duplicate_columns",
        "duplicate_row_vector",
        "fixed_ratio",
        "fixed_difference",
        "row_offset_exact_reuse",
        "paired_difference_spread",
        "cross_sheet_duplication",
        "p_value_inconsistency",
    ]
    generated = [
        _generate_one(root, twins_root, klass, seq)
        for klass in classes
        for seq in range(1, 4)
    ]
    summary = {
        "paper_id": PAPER_ID,
        "base_dir": str(root),
        "twins_root": str(twins_root),
        "cleanliness": _source_data_cleanliness(root),
        "statcheck_clean_findings": _statcheck_probe(root),
        "grim_probe_n2249_dec14": {
            "status": "removed_from_cns098_v1_1",
            "reason": "Mean logHR is continuous; GRIM moved to P1 substrate.",
        },
        "numeric_summary_tables": {
            path.name: [table.__dict__ for table in discover_numeric_summary_tables(path)]
            for path in sorted(root.glob("*.xlsx"))
        },
        "generated": generated,
    }
    output = root / "round1_summary.json"
    output.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "twins": len(generated)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
