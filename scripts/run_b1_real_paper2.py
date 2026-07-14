"""B1 bare agent, REAL backbones, on REAL fabrication data (paper2, NPC) with real sheet content.

Each claim asks: "is this source-data sheet genuine independent measurements?" The content
renderer loads the actual xlsx rows so the model can spot duplicated rows / fixed ratios itself.
Dirty sheets (PubPeer-confirmed) should be judged inconsistent; benign sheets consistent. This is
the first real look at how bare competitor agents do on real fabrication. Reads keys from .env.

Run:  VERITAS_REAL_PAPERS_ROOT=<root> PYTHONPATH=. python3 scripts/run_b1_real_paper2.py
"""

from __future__ import annotations

import sys

from openpyxl import load_workbook

from engine.benchmark.agent_harness import run_bare_agent
from engine.benchmark.backbones import make_json_agent
from engine.benchmark.metrics import main_table_row
from engine.benchmark.real_paper import case_data_dir
from engine.benchmark.schema import (
    LABEL_CLEAN,
    LABEL_DIRTY,
    LEVEL_RELATION,
    BenchmarkCase,
    ClaimInstance,
    Evidence,
)
from engine.benchmark.verifiers import source_data_renderer
from scripts.live_backbones import AUDITOR_SYSTEM, CANONICAL_MODELS, make_backbone

DOI = "10.1038/s41588-025-02253-8"
CLAIM_TEXT = ("This source-data sheet reports genuine, independently-measured values — there is no "
              "internal duplication, copied row-block, or fixed ratio/difference between rows or columns.")
DIRTY = [
    ("41588_2025_2253_MOESM8_ESM.xlsx", "Fig.5c"), ("41588_2025_2253_MOESM9_ESM.xlsx", "Fig.6i"),
    ("41588_2025_2253_MOESM11_ESM.xlsx", "Fig.8f"), ("41588_2025_2253_MOESM6_ESM.xlsx", "Fig.3f"),
    ("41588_2025_2253_MOESM10_ESM.xlsx", "Fig.7d"), ("41588_2025_2253_MOESM13_ESM.xlsx", "Extended Data Fig.3g"),
]
KEYS = ["claim_f1", "claim_recall", "far", "coverage", "abstention"]


def _claim(workbook, sheet, label):
    return ClaimInstance(
        claim_id=f"{workbook}:{sheet}", claim_type="source_data.duplicate_columns", level="L3",
        label=label, relation=LEVEL_RELATION["L3"], claim_text=CLAIM_TEXT,
        evidence=Evidence(evidence_type="source_data", target=f"{workbook} / {sheet}", workbook=workbook, sheet=sheet),
        evidence_span=f"L3:{workbook}->{workbook}:{sheet}")


def _benign(data_dir, dirty, limit=6):
    out = []
    for wb_path in sorted(data_dir.glob("*.xlsx")):
        wb = load_workbook(wb_path, read_only=True)
        for ws in wb.worksheets:
            if (wb_path.name, ws.title) in dirty or (ws.max_row or 0) < 8:
                continue
            out.append(_claim(wb_path.name, ws.title, LABEL_CLEAN))
            if len(out) >= limit:
                wb.close()
                return out
        wb.close()
    return out


def main() -> int:
    data_dir = case_data_dir(DOI)
    if not data_dir.exists():
        print(f"missing data at {data_dir}; set VERITAS_REAL_PAPERS_ROOT", file=sys.stderr)
        return 1
    dirty_set = set(DIRTY)
    claims = [_claim(wb, sh, LABEL_DIRTY) for wb, sh in DIRTY] + _benign(data_dir, dirty_set)
    case = BenchmarkCase("NPC-paper2", "real-test", DOI, tuple(claims))
    render = source_data_renderer(data_dir)
    n_dirty = sum(c.label == LABEL_DIRTY for c in claims)
    n_clean = sum(c.label == LABEL_CLEAN for c in claims)
    print(f"paper2 REAL: {n_dirty} dirty + {n_clean} benign sheets, real content, live B1 agents\n")
    print(f"{'backbone':18s} " + " ".join(f"{k:>10s}" for k in KEYS) + "   dirty_caught")
    for model in CANONICAL_MODELS:
        try:
            agent = make_json_agent(make_backbone(model), system=AUDITOR_SYSTEM)
            preds = run_bare_agent(case, agent, render=render)
            row = main_table_row(preds)
            caught = sum(1 for p in preds if p.gt_label == LABEL_DIRTY and p.predicted_flag)
            print(f"{model:18s} " + " ".join(f"{row[k]:>10.3f}" for k in KEYS) + f"   {caught}/{n_dirty}")
        except Exception as e:  # noqa: BLE001
            print(f"{model:18s} ERROR: {type(e).__name__}: {str(e)[:70]}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
